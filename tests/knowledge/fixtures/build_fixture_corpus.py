"""Build the synthetic fixture corpus shared by every knowledge-base test (handoff KWP1).

Everything here is invented: the vendors (Coastline Home Inspections, Tidewater
Plumbing & Rooter, Bayside Electric Co., Harborline), the people, the address,
the amounts and every reference number. No house data (rule K2).

Run it to regenerate ``tests/knowledge/fixtures/files/``::

    uv run python -m tests.knowledge.fixtures.build_fixture_corpus

The output is byte-for-byte reproducible for the locked ``reportlab`` and
``python-docx`` versions (``invariant=1`` PDFs with uncompressed page streams,
a DOCX with fixed timestamps and member order); a test
rebuilds it into a temp dir and compares. :mod:`tests.knowledge.fixtures.corpus`
exposes the paths and the ground truth that later work packages test against.

Files:

- ``inspection.pdf``: 12 pages. A DocuSign-style envelope line and an
  inspector line on top of every page, the property address and "Page N of 12"
  at the bottom, a summary list, three sections of numbered findings with
  severity and recommendation, a photo page with **no text layer** (p. 9), a
  hyphenated standards page, an invoice page and a certification page.
- ``estimate.pdf``: 2 pages with letter-spaced headings ("D ES C R I PT I O N O F").
- ``invoice.pdf``: 1 page.
- ``bundle.pdf``: the three above concatenated (15 pages), for segmentation.
- ``unnumbered_bundle.pdf``: 5 pages and 4 documents with no page numbering at
  all: two back-to-back invoices from the same vendor (the second without an
  invoice number), a 2-page roofing proposal whose second page has a
  "Roofing Services Overview" heading but no letterhead, and a pest-control
  completion notice.
- ``mixed_bundle.pdf``: 22 pages and 10 documents in the shapes a real closing
  bundle takes: a 2-page form numbered "(FORM PAGE 1 OF 2)"; a 4-page
  inspection whose letterhead line ends in "Page k of 4", with "Moderate" and
  "Recommendation" labels that also sit at a fixed bottom position; an invoice
  and two estimates from one vendor with "Invoice #… Page k of N" footers; two
  pest reports numbered only by a bare "k of N" at the end of a form row; a
  completion notice with a long title; a letter-style invoice whose top lines
  are a letterhead (address, licence, phone, date); and an invoice whose
  contact block is at the bottom. A DocuSign line runs along pp. 3-17.
- ``review_b1.pdf``, ``review_b1b.pdf``, ``review_b2.pdf``, ``review_b3.pdf``:
  the KWP1b reviewer's over-fitting probes (manuals with step and figure
  counters, letters, logs, permits, invoices, notices; see ``review_bundle1``).
- ``heat_pump_guide.docx``: a title and three top-level headings.
- ``irrigation_notes.html``: a title and three second-level sections, plus
  script and style that must not leak into the text.
- ``garage_door.md``: a title and three ``##`` sections, with a fenced block.
- ``pool_log.txt``: about 7,000 characters with no headings.
"""

from __future__ import annotations

import argparse
import io
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

FILES_DIR = Path(__file__).resolve().parent / "files"

# ------------------------------------------------------------------ content
ENVELOPE_LINE = "Docusign Envelope ID: 5E7C1A90-3B2D-4C6E-9F10-A1B2C3D4E5F6"
INSPECTOR_LINE = "Coastline Home Inspections | Inspector: Jordan Avery Reyes | Report CHI-2025-0314"
ADDRESS = "742 Driftwood Way, Port Alder, CA 90000"
INSPECTION_PAGES = 12
INSPECTION_VENDOR = "Coastline Home Inspections"
INSPECTION_DATE = "2025-03-14"
ESTIMATE_VENDOR = "Tidewater Plumbing & Rooter, Inc."
ESTIMATE_DATE = "2025-03-20"
INVOICE_VENDOR = "Bayside Electric Co."
INVOICE_DATE = "2025-04-02"
PHOTO_PAGE = 9

#: The recommendation line that repeats on exactly two pages (pp. 3 and 4), at
#: the bottom of each: cleaning must keep it.
REPEATED_LINE = "Recommendation: Recommend a qualified contractor evaluate and repair."

#: Letter-spaced headings exactly as PDFium extracts them, and their repair.
LETTER_SPACED = {
    "E S T I M A T E": "ESTIMATE",
    "D ES C R I PT I O N O F W O R K": "DESCRIPTION OF WORK",
    "T E R M S A N D C O N D I T I O N S": "TERMS AND CONDITIONS",
}


@dataclass(frozen=True)
class Finding:
    id: str
    area: str
    component: str
    title: str
    severity: str
    description: str
    recommendation: str
    page: int

    @property
    def heading(self) -> str:
        return f"{self.id} {self.area} – {self.component}: {self.title}"


SECTIONS = {
    1: ("ROOF AND EXTERIOR", (3, 4)),
    2: ("ELECTRICAL", (5, 6)),
    3: ("PLUMBING", (7, 8)),
}

FINDINGS: tuple[Finding, ...] = (
    Finding("1.1.1", "Roof", "Covering", "Cracked Tiles", "Repair",
            "Four cracked concrete tiles were observed on the south slope above the garage. "
            "Cracked tiles can let water reach the underlayment and the roof sheathing.",
            "Recommendation: Recommend a licensed roofing contractor replace the cracked tiles.", 3),
    Finding("1.1.2", "Roof", "Flashing", "Loose Step Flashing", "Repair",
            "Step flashing at the chimney is loose and the sealant along its top edge has "
            "failed. Staining on the attic sheathing below suggests past leaks.",
            REPEATED_LINE, 3),
    Finding("1.2.1", "Exterior", "Gutters", "Downspout Discharges at Foundation", "Maintenance",
            "The rear left downspout discharges next to the foundation, and the soil there "
            "slopes toward the house.",
            "Recommendation: Recommend extending the downspout at least 4 feet from the house.", 4),
    Finding("1.3.1", "Exterior", "Siding", "Wood Rot at Fascia Trim", "Repair",
            "Soft, deteriorated wood was found at the fascia trim near the back door. The "
            "paint has peeled and the wood is exposed.",
            REPEATED_LINE, 4),
    Finding("2.1.1", "Electrical", "Service Panel", "Double-Tapped Breaker", "Safety Hazard",
            "Two conductors are connected to a single-pole 20-amp breaker in the main panel. "
            "This breaker is not rated for two conductors.",
            "Recommendation: Recommend a licensed electrician correct the double tap.", 5),
    Finding("2.2.1", "Electrical", "Receptacles", "No GFCI Protection in Kitchen",
            "Safety Hazard",
            "Kitchen counter receptacles to the right of the sink are not GFCI protected.",
            "Recommendation: Recommend a licensed electrician install GFCI protection.", 5),
    Finding("2.3.1", "Electrical", "Smoke Alarms", "Missing in Hallway", "Safety Hazard",
            "No smoke alarm was found in the hallway outside the bedrooms.",
            "Recommendation: Recommend installing a smoke alarm in the hallway.", 6),
    Finding("2.3.2", "Electrical", "Lighting", "Inoperable Garage Light", "Monitor",
            "The garage ceiling light did not respond to the switch. The bulb may be "
            "burned out.",
            "Recommendation: Recommend replacing the bulb and testing the fixture again.", 6),
    Finding("3.1.1", "Plumbing", "Water Heater", "Age", "Monitor",
            "The 50-gallon gas water heater was manufactured in 2012 and is 13 years old, "
            "near the end of its typical service life.",
            "Recommendation: Recommend budgeting for replacement.", 7),
    Finding("3.1.2", "Plumbing", "Water Heater", "Valve Leak", "Repair",
            "The temperature and pressure relief valve is dripping, and mineral deposits "
            "have formed at the discharge pipe.",
            "Recommendation: Recommend a licensed plumber replace the relief valve.", 7),
    Finding("3.1.3", "Plumbing", "Water Heater", "No Expansion Tank", "Maintenance",
            "No thermal expansion tank is installed on the cold water supply line.",
            "Recommendation: Recommend a licensed plumber install a thermal expansion tank.", 7),
    Finding("3.2.1", "Plumbing", "Supply Lines", "Galvanized Piping", "Monitor",
            "Galvanized steel supply piping was observed in the crawlspace. Reduced water "
            "pressure was noted at the upstairs shower.",
            "Recommendation: Recommend a licensed plumber evaluate the supply piping.", 8),
    Finding("3.3.1", "Plumbing", "Main Shutoff", "Location", "Maintenance",
            "The main water shutoff valve is located at the front of the house, left of the "
            "garage door.",
            "Recommendation: Recommend labeling the valve so it can be found quickly.", 8),
)  # fmt: skip

INVOICE_PAGE_LINES = (
    "Coastline Home Inspections",
    "Invoice number: CHI-INV-20250314",
    "Invoice date: March 14, 2025",
    "Bill to: Morgan Example",
    f"Service address: {ADDRESS}",
    "",
    "Home inspection, single-family (1,850 sq ft): $475.00",
    "Sewer lateral video scope: $150.00",
    "Total due: $625.00",
    "",
    "Paid in full on March 14, 2025. Balance due: $0.00",
    "Thank you for choosing Coastline Home Inspections.",
)

STANDARDS_LINES = (
    "This inspection was performed in accordance with the Standards of Practice of the",
    "Pacific Association of Home Inspectors (a fictional association). The inspection is a",
    "visual, non-invasive examination of the readily accessible systems and components of the",
    "home. The inspector is not required to deter-",
    "mine the remaining life of any system or component, to move furniture or stored items,",
    "or to enter any area that is unsafe or not readily accessible.",
    "",
    "The report describes the condition of the house on the day of the inspection. Conditions",
    "change over time, and items that were working on the day of the inspection may fail",
    "later. The client should have every item marked Repair or Safety Hazard evaluated by a",
    "licensed profes-",
    "sional before the end of the contingency period.",
)


# ------------------------------------------------------------------ drawing
PAGE_W, PAGE_H = 612, 792
LEFT, RIGHT = 60, 552
BODY_FONT, BOLD_FONT = "Helvetica", "Helvetica-Bold"


class _Page:
    """A cursor over one reportlab page."""

    def __init__(self, canvas, top: float = 720) -> None:
        self.c = canvas
        self.y = top

    def line(self, text: str, font: str = BODY_FONT, size: float = 10, lead: float = 14) -> None:
        if text:
            self.c.setFont(font, size)
            self.c.drawString(LEFT, self.y, text)
        self.y -= lead

    def para(self, text: str, font: str = BODY_FONT, size: float = 10, lead: float = 13) -> None:
        from reportlab.lib.utils import simpleSplit

        for part in simpleSplit(text, font, size, RIGHT - LEFT):
            self.line(part, font, size, lead)

    def gap(self, amount: float = 8) -> None:
        self.y -= amount

    def spaced(self, spaced: str, size: float = 13) -> None:
        """Draw a heading glyph group by glyph group, as a letter-spaced PDF does."""
        from reportlab.pdfbase.pdfmetrics import stringWidth

        self.c.setFont(BOLD_FONT, size)
        x = LEFT
        for group in spaced.split(" "):
            self.c.drawString(x, self.y, group)
            x += stringWidth(group, BOLD_FONT, size) + 6
        self.y -= size + 8


def _inspection_header(canvas) -> None:
    canvas.setFont(BODY_FONT, 7)
    canvas.drawString(LEFT, 772, ENVELOPE_LINE)
    canvas.drawString(LEFT, 760, INSPECTOR_LINE)


def _inspection_footer(canvas, number: int) -> None:
    canvas.setFont(BODY_FONT, 7)
    canvas.drawString(LEFT, 40, ADDRESS)
    canvas.drawRightString(RIGHT, 28, f"Page {number} of {INSPECTION_PAGES}")


def _photo_page(canvas) -> None:
    """A page that is only a picture: no text at all, so no text layer."""
    from PIL import Image, ImageDraw
    from reportlab.lib.utils import ImageReader

    image = Image.new("RGB", (160, 120))
    pixels = image.load()
    for x in range(160):
        for y in range(120):
            pixels[x, y] = (40 + x // 2, 90 + y, 150 - x // 3)
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 30, 70, 90), fill=(170, 60, 40))
    draw.polygon([(90, 90), (120, 40), (150, 90)], fill=(90, 90, 90))
    canvas.drawImage(ImageReader(image), 96, 250, width=420, height=315)


def _findings_page(canvas, number: int, section: int, first_page: bool) -> None:
    title, _ = SECTIONS[section]
    page = _Page(canvas)
    if first_page:
        page.line(f"{section}. {title}", BOLD_FONT, 14, 22)
    else:
        page.line(f"{section}. {title} (continued)", BOLD_FONT, 12, 20)
    for finding in (f for f in FINDINGS if f.page == number):
        page.line(finding.heading, BOLD_FONT, 11, 15)
        page.line(f"Severity: {finding.severity}", BODY_FONT, 10, 14)
        page.para(finding.description)
        page.line(finding.recommendation)
        page.gap(10)


def draw_inspection(canvas) -> None:
    for number in range(1, INSPECTION_PAGES + 1):
        if number == PHOTO_PAGE:
            _photo_page(canvas)
            canvas.showPage()
            continue
        _inspection_header(canvas)
        page = _Page(canvas)
        if number == 1:
            page.line("HOME INSPECTION REPORT", BOLD_FONT, 20, 30)
            for text in (
                f"Property: {ADDRESS}",
                "Client: Morgan Example",
                "Inspection date: March 14, 2025",
                "Inspector: Jordan Avery Reyes, License HI-000000",
                "Prepared by Coastline Home Inspections",
                "Report number: CHI-2025-0314",
                "",
                "Weather: Clear, 64 degrees F. The house was vacant and the utilities were on.",
                "Year built: 1987. Approximate size: 1,850 sq ft. Foundation: raised, with crawlspace.",
            ):
                page.line(text, lead=16)
        elif number == 2:
            page.line("SUMMARY", BOLD_FONT, 14, 22)
            page.para(
                "This summary lists the items that need attention. It is not the whole report; "
                "read every section, including the limitations."
            )
            page.gap()
            for finding in FINDINGS:
                page.line(f"{finding.heading} ({finding.severity})", lead=15)
        elif number in (3, 5, 7):
            _findings_page(canvas, number, (number - 1) // 2, first_page=True)
        elif number in (4, 6, 8):
            _findings_page(canvas, number, (number - 2) // 2, first_page=False)
        elif number == 10:
            page.line("STANDARDS OF PRACTICE AND LIMITATIONS", BOLD_FONT, 13, 20)
            for text in STANDARDS_LINES:
                page.line(text, lead=13)
        elif number == 11:
            page.line("INVOICE", BOLD_FONT, 16, 24)
            for text in INVOICE_PAGE_LINES:
                page.line(text, lead=15)
        elif number == 12:
            page.line("INSPECTOR CERTIFICATION", BOLD_FONT, 13, 20)
            page.para(
                "I certify that I personally inspected the property named in this report and "
                "that I have no interest in the property or in any repairs. The report was "
                "completed on March 15, 2025."
            )
            page.gap()
            page.line("Jordan Avery Reyes, Coastline Home Inspections")
        _inspection_footer(canvas, number)
        canvas.showPage()


ESTIMATE_LETTERHEAD = (ESTIMATE_VENDOR, "License C36-000000 | (555) 010-0100 | Port Alder, CA")


def draw_estimate(canvas) -> None:
    # page 1
    page = _Page(canvas, top=740)
    page.line(ESTIMATE_LETTERHEAD[0], BOLD_FONT, 15, 18)
    page.line(ESTIMATE_LETTERHEAD[1], BODY_FONT, 8, 26)
    page.spaced("E S T I M A T E", 16)
    for text in (
        "Estimate No. TPR-58213",
        "Date: March 20, 2025",
        "Customer: Morgan Example",
        f"Job site: {ADDRESS}",
    ):
        page.line(text)
    page.gap(10)
    page.spaced("D ES C R I PT I O N O F W O R K")
    for text in (
        "1. Replace the water heater temperature and pressure relief valve: $185.00",
        "2. Install a 2-gallon thermal expansion tank on the cold water line: $420.00",
        "",
        "Subtotal: $605.00",
        "Tax: $0.00",
        "Total: $605.00",
    ):
        page.line(text)
    canvas.setFont(BODY_FONT, 8)
    canvas.drawRightString(RIGHT, 28, "Page 1 of 2")
    canvas.showPage()
    # page 2
    page = _Page(canvas, top=740)
    page.line(ESTIMATE_LETTERHEAD[0], BOLD_FONT, 15, 18)
    page.line(ESTIMATE_LETTERHEAD[1], BODY_FONT, 8, 26)
    page.spaced("T E R M S A N D C O N D I T I O N S")
    page.para(
        "This estimate is valid for 30 days. A 50% deposit is due before work begins and the "
        "balance is due on completion. All work carries a one-year labor warranty. Permits, "
        "if required, are included in the price."
    )
    page.gap(20)
    page.line("Accepted by: ______________________    Date: ____________")
    canvas.setFont(BODY_FONT, 8)
    canvas.drawRightString(RIGHT, 28, "Page 2 of 2")
    canvas.showPage()


def draw_invoice(canvas) -> None:
    page = _Page(canvas, top=740)
    page.line(INVOICE_VENDOR, BOLD_FONT, 15, 18)
    page.line("License C10-000000 | (555) 010-0199", BODY_FONT, 8, 26)
    page.line("INVOICE", BOLD_FONT, 16, 24)
    for text in (
        "Invoice # BE-10442",
        "Date: April 2, 2025",
        "Customer: Morgan Example",
        f"Service address: {ADDRESS}",
    ):
        page.line(text)
    page.gap()
    page.para(
        "Work performed: Separated the double-tapped conductors in the main panel and "
        "installed a new 20-amp breaker. Installed a GFCI receptacle at the kitchen counter."
    )
    page.gap()
    for text in ("Labor (2.0 hours): $240.00", "Materials: $100.00", "Total: $340.00"):
        page.line(text)
    canvas.setFont(BODY_FONT, 8)
    canvas.drawRightString(RIGHT, 28, "Page 1 of 1")
    canvas.showPage()


UNNUMBERED_VENDORS = (
    "Driftwood Pool Service LLC",
    "Harbor Point Roofing Co.",
    "Pelican Pest Control, Inc.",
)
UNNUMBERED_DATES = ("2025-05-05", "2025-06-05", "2025-05-12", "2025-06-20")


def draw_unnumbered(canvas) -> None:
    """Four documents on five pages, with no "Page N of M" anywhere."""
    pool, roofing, pest = UNNUMBERED_VENDORS
    pages = (
        (pool, "INVOICE", ("Invoice # DPS-1001", "Date: May 5, 2025", "Customer: Morgan Example",
                           "Monthly pool service (May): $140.00", "Total: $140.00")),
        (pool, "INVOICE", ("Date: June 5, 2025", "Customer: Morgan Example",
                           "Monthly pool service (June): $140.00",
                           "Replacement filter cartridge: $65.00", "Total: $205.00")),
        (roofing, "PROPOSAL", ("Proposal No. HPR-311", "Date: May 12, 2025",
                               f"Job site: {ADDRESS}",
                               "Scope: replace cracked concrete tiles on the south slope.")),
        (None, None, ("Scope of work (continued)", "Reseal the chimney step flashing.",
                      "Roofing Services Overview",
                      "All work is done by our own crews and carries a 5-year warranty.",
                      "Proposal total: $1,480.00")),
        (pest, "NOTICE OF WORK COMPLETED", ("Date: June 20, 2025", f"Property: {ADDRESS}",
                                            "Fumigation for drywood termites was completed.")),
    )  # fmt: skip
    for vendor, form, lines in pages:
        page = _Page(canvas, top=740)
        if vendor:
            page.line(vendor, BOLD_FONT, 15, 22)
        if form:
            page.line(form, BOLD_FONT, 16, 24)
        for text in lines:
            page.line(text)
        canvas.showPage()


MIXED_STARTS = (1, 3, 7, 10, 12, 14, 18, 19, 21, 22)
MIXED_TYPES = (
    "disclosure_form", "inspection_report", "invoice", "estimate", "estimate",
    "pest_report", "completion_notice", "pest_report", "invoice", "invoice",
)  # fmt: skip
MIXED_ENVELOPE = "Docusign Envelope ID: 3C9A51E2-6B7D-4F80-9A1B-2C3D4E5F6A7B"
MIXED_WDO_LABELS = "BLDG NO. STREET CITY ZIP INSPECTION DATE NUMBER OF PAGES"
_OSPREY = ("Osprey Plumbing & Drain", "License #000321 | Port Alder, CA 90000")


def _lines_page(canvas, top: list[str], bottom: list[str] = ()) -> None:
    """Draw ``top`` from the top down, then ``bottom`` upward from the foot (drawn last)."""
    page = _Page(canvas, top=750)
    for text in top:
        page.line(text, lead=14)
    canvas.setFont(BODY_FONT, 8)
    for index, text in enumerate(bottom):  # in reading order, so the stream order matches
        canvas.drawString(LEFT, 28 + 12 * (len(bottom) - 1 - index), text)
    canvas.showPage()


def _mixed_pages() -> list[tuple[list[str], list[str]]]:
    env = [MIXED_ENVELOPE]
    pages: list[tuple[list[str], list[str]]] = []
    # D1: a two-page disclosure form
    for k in (1, 2):
        pages.append(
            (
                [
                    "RECEIPT FOR REPORTS",
                    f"(FORM RR PAGE {k} OF 2)",
                    "The buyer acknowledges receipt of the reports listed below.",
                    *(f"Report line {k}.{n}: received and reviewed." for n in range(1, 5)),
                ],
                ["Form RR, a fictional disclosure form"],
            )
        )
    # D2: a four-page inspection, letterhead line carries the page number
    for k in range(1, 5):
        top = [
            f"Harbor Point Home Inspections Page {k} of 4",
            "(555) 010-0300",
            "office@harborpoint.example",
        ]
        if k == 1:
            top.append("PROPERTY INSPECTION REPORT")
        body = [
            f"{k}.1 Finding about the roof, part {k}",
            "Moderate",
            f"The flashing at corner {k} is loose.",
            "Recommendation",
            "Have a roofer reseal it.",
            f"{k}.2 Finding about the gutters, part {k}",
            "Moderate",
            f"Gutter run {k} sags.",
            "Recommendation",
            "Rehang the gutter.",
        ]
        pages.append((top + body, env))
    # D3-D5: one vendor's invoice and two estimates, numbered in the footer
    for kind, number, count, date_line in (
        ("INVOICE", "40010001", 3, "May 2, 2025"),
        ("ESTIMATE", "40010002", 2, "May 2, 2025"),
        ("ESTIMATE", "40010003", 2, "May 2, 2025"),
    ):
        for k in range(1, count + 1):
            top = [kind, number, f"{kind} DATE", date_line, *_OSPREY] if k == 1 else []
            top += [f"{kind.title()} line {k}.{n}: pipe work item {n}" for n in range(1, 4)]
            pages.append((top, [f"{kind.title()} #{number} Page {k} of {count}", *env]))
    # D6: a four-page pest report, numbered only by "k of 4" in a form row
    values = "12 GULL LN PORT ALDER 90000 04/02/2025"
    for k in range(1, 5):
        if k == 1:
            top = [
                "Inspector / Licence No. 000456",
                "Sandpiper Pest Control, Inc.",
                "WOOD DESTROYING PESTS AND ORGANISMS INSPECTION REPORT",
                "Findings are listed by section below.",
            ]
            bottom = [MIXED_WDO_LABELS, f"{values} {k} of 4", *env]
        else:
            top = [MIXED_WDO_LABELS, f"{values} {k} of 4", f"Section {k}: subarea findings"]
            bottom = ["Form: WDO-7 -- Report #: 70001", *env]
        pages.append((top + [f"Item {k}A: evidence of drywood termites at the eave."], bottom))
    # D7: a completion notice with a long title
    pages.append(
        (
            [
                "Standardized Notice of Work Completed; Items Pending",
                "Report #: 70001",
                "Items 1A and 2A were completed on April 20, 2025.",
            ],
            ["Sandpiper Pest Control, Inc.", "(555) 010-0400 - Phone"],
        )
    )
    # D8: a two-page reinspection, "k of 2" in the form row
    values = "12 GULL LN PORT ALDER 90000 05/01/2025"
    for k in (1, 2):
        top = (
            [
                "Inspector / Licence No. 000456",
                "Sandpiper Pest Control, Inc.",
                "WOOD DESTROYING PESTS AND ORGANISMS REINSPECTION REPORT",
            ]
            if k == 1
            else [MIXED_WDO_LABELS, f"{values} {k} of 2"]
        )
        bottom = [MIXED_WDO_LABELS, f"{values} {k} of 2"] if k == 1 else ["Report #: 70002"]
        pages.append((top + [f"Reinspection note {k}: no further activity."], bottom))
    # D9: a letter-style invoice: the top lines are a letterhead
    pages.append(
        (
            [
                "Marisol Vega Finish Carpentry",
                "88 Kelp Ave.",
                "Port Alder, CA. 90000",
                "Lic # 000111",
                "(555) 010-0444",
                "June 3, 2025",
                "Re: 12 Gull Ln",
                "Repaired the eave boards noted in the pest report.",
                "Total: $900",
            ],
            [],
        )
    )
    # D10: an invoice whose contact block is at the bottom
    pages.append(
        (
            [
                "Pay online",
                "ITEM QTY AMOUNT",
                "Thermostat, programmable 1 $180.00",
                "Invoice # GA-17-2",
                "Total: $180.00",
            ],
            [
                "Gull Air Heating",
                "Questions? Call (555) 010-0555 billing@gullair.example",
                "License #000222",
            ],
        )
    )
    return pages


def draw_mixed(canvas) -> None:
    for top, bottom in _mixed_pages():
        _lines_page(canvas, top, bottom)


# ------------------------------------------- review bundles (KWP1b review, invented)
# Built by the KWP1b reviewer to catch over-fitted segmentation rules: manuals with
# "Step 3 of 5" and "Figure k of 3" lines, letters whose last page names another
# phone number, dated logs and journals, permits with dated inspection records.
REVIEW_FILL = [
    "The unit should be inspected at least once a year by a qualified technician.",
    "Keep the area around the equipment clear of debris and stored items.",
    "Refer to the table below for recommended settings in each season.",
    "If the indicator light flashes, turn off power and wait thirty seconds.",
    "Clean the filter screen with warm water and a soft brush.",
    "Do not use abrasive cleaners on the stainless surfaces.",
    "Record each service visit so the history is available later.",
    "Normal operating noise includes a soft hum and occasional clicks.",
    "Contact a licensed professional for any electrical repair.",
    "Replace worn gaskets before they start to leak.",
    "The drain line must slope continuously toward the outlet.",
    "Allow the system to run a full cycle after any adjustment.",
]


def _fill(tag, n=10, extra=()):
    lines = [f"{tag}: {REVIEW_FILL[i % len(REVIEW_FILL)]}" for i in range(n)]
    return list(extra) + lines


def _draw_review(canvas, pages):
    c = canvas
    for top, mid, bottom in pages:
        y = 750
        for t in top + mid:
            c.setFont("Helvetica", 10)
            c.drawString(54, y, t)
            y -= 14
        for i, t in enumerate(bottom):
            c.setFont("Helvetica", 8)
            c.drawString(54, 28 + 12 * (len(bottom) - 1 - i), t)
        c.showPage()


def review_bundle1(step=True):
    docs = []
    # D1 dishwasher manual, 6 pages, running header, bare page-number footer
    pages = []
    for k in range(1, 7):
        top = ["AquaSteel DW-500 Dishwasher", "Use and Care Guide"]
        if k == 1:
            top += ["OWNERS MANUAL", "Rev. 03/2019"]
        mid = _fill(f"Section {k}", 10)
        if k == 5:
            mid[4] = "LIMITED WARRANTY"
        bottom = []
        if k == 1:
            bottom.append("Customer Care: 1-800-555-0100  www.aquasteel.example")
        if k == 4 and step:
            mid[-1] = "Installation, Step 3 of 5"
        if k == 6:
            bottom.append("Parts and accessories: 1-800-555-0177")
        bottom.append(str(k))
        pages.append((top, mid, bottom))
    docs.append(("manual", pages))
    # D2 warranty registration card, same phone
    docs.append(
        (
            "warranty",
            [
                (
                    [
                        "AquaSteel Appliances Inc.",
                        "PRODUCT WARRANTY REGISTRATION",
                        "Call 1-800-555-0100 to register by phone",
                    ],
                    _fill("Card", 8),
                    [],
                )
            ],
        )
    )
    # D3 HOA letter, 2 pages, signed by the management company
    p1 = (
        [
            "Coastal Bluffs Homeowners Association",
            "PO Box 100, Seaview, CA 92000",
            "(760) 555-0101 | www.coastalbluffs.example",
            "March 3, 2024",
            "Dear Homeowner,",
        ],
        _fill("Letter", 10),
        [],
    )
    p2 = (
        ["Homeowner, 12 Example Lane", "March 3, 2024", "Page 2"],
        _fill("Letter cont", 8),
        ["Sincerely, the Board", "Managed by Sunset Property Management, Inc. (760) 555-0199"],
    )
    docs.append(("other", [p1, p2]))
    # D4 building permit, 3 pages, inspection record with dates at top
    docs.append(
        (
            "permit",
            [
                (
                    [
                        "City of Seaview Development Services",
                        "BUILDING PERMIT",
                        "Permit No. B24-0117",
                        "Issued: 04/02/2024",
                    ],
                    _fill("Scope", 10),
                    ["City of Seaview (760) 555-0300"],
                ),
                (
                    [
                        "Inspection Record - Permit B24-0117",
                        "Footing 04/10/2024 Approved",
                        "Framing 04/20/2024 Approved",
                        "Rough electrical 04/22/2024 Approved",
                    ],
                    _fill("Inspection notes", 9),
                    ["City of Seaview (760) 555-0300"],
                ),
                (
                    ["Final 05/15/2024 Approved", "Conditions of approval"],
                    _fill("Conditions", 10),
                    ["City of Seaview (760) 555-0300"],
                ),
            ],
        )
    )
    # D5, D6 two invoices from one vendor
    for num, day in (("5521", "06/03/2024"), ("5534", "06/10/2024")):
        docs.append(
            (
                "invoice",
                [
                    (
                        [
                            "Bayside Electric Co.",
                            "Lic. No. 998877  (760) 555-0400",
                            "INVOICE",
                            f"Invoice # {num}",
                            f"Date: {day}",
                        ],
                        _fill("Line item", 8),
                        ["Thank you for your business"],
                    )
                ],
            )
        )
    # D7 pool service log, 4 pages, date at top of each page
    pages = []
    for k, wk in enumerate(("2024-05-06", "2024-05-13", "2024-05-20", "2024-05-27"), 1):
        pages.append(
            (
                ["Blue Wave Pool Service - Weekly Service Log", f"Week of {wk}"],
                _fill(f"Visit {k}", 10),
                ["Blue Wave Pool Service (760) 555-0500"],
            )
        )
    docs.append(("other", pages))
    # D8 roof warranty, 2 pages, second page heading repeats the form word
    docs.append(
        (
            "warranty",
            [
                (
                    ["Summit Roofing, Inc.", "LIMITED WORKMANSHIP WARRANTY", "(760) 555-0600"],
                    _fill("Coverage", 10),
                    [],
                ),
                (["WARRANTY TERMS AND CONDITIONS (continued)"], _fill("Terms", 11), []),
            ],
        )
    )
    return docs


def review_bundle2():
    docs = []
    # D1 pool pump manual, 8 pages, figure captions "Figure k of 3" near the foot
    pages = []
    for k in range(1, 9):
        top = ["HydroMax Variable Speed Pump", "Installation and User Guide"]
        mid = _fill(f"Chapter {k}", 11)
        bottom = []
        if k == 1:
            bottom.append("Technical support (800) 555-0142")
        if k in (3, 4):
            bottom.append(f"Wiring diagram, Figure {k - 2} of 3")
        if k == 7:
            bottom.append("Still stuck? Call our service line (888) 555-0190")
        bottom.append(f"- {k} -")
        pages.append((top, mid, bottom))
    docs.append(("manual", pages))
    # D2, D3 two monthly statements from the utility, "Page k of 2"
    for month in ("01/15/2024", "02/15/2024"):
        pages = []
        for k in (1, 2):
            top = (
                ["Seaview Water District", f"Statement date {month}"]
                if k == 1
                else ["Seaview Water District", "Usage detail"]
            )
            pages.append((top, _fill("Usage", 10), [f"Page {k} of 2", "(760) 555-0700"]))
        docs.append(("other", pages))
    # D4 one-page letter from a tile contractor
    docs.append(
        (
            "other",
            [
                (
                    ["Rivera Tile and Stone", "Lic # 445566", "(760) 555-0800", "Dear Owner,"],
                    _fill("Tile note", 8),
                    [],
                )
            ],
        )
    )
    # D5 renovation journal, 3 pages, a dated heading at the top of each page
    pages = []
    for d in ("Tuesday, June 4, 2024", "Friday, June 7, 2024", "Monday, June 10, 2024"):
        pages.append(([d, "Kitchen remodel journal"], _fill("Journal", 11), []))
    docs.append(("other", pages))
    # D6 HOA violation notice
    docs.append(
        (
            "other",
            [
                (
                    [
                        "Coastal Bluffs Homeowners Association",
                        "August 1, 2024",
                        "NOTICE OF VIOLATION",
                        "(760) 555-0101",
                    ],
                    _fill("Notice", 8),
                    [],
                )
            ],
        )
    )
    # D7 two-page quote
    docs.append(
        (
            "estimate",
            [
                (
                    ["Pacific Solar Screens LLC", "QUOTE", "Quote # Q-778", "July 2, 2024"],
                    _fill("Quote item", 9),
                    ["(760) 555-0900"],
                ),
                (
                    ["Pacific Solar Screens LLC", "Quote # Q-778 (continued)"],
                    _fill("Quote terms", 9),
                    ["(760) 555-0900"],
                ),
            ],
        )
    )
    return docs


def review_bundle3():
    docs = []
    # D1 3-page report, "Page k of 3"
    docs.append(
        (
            "inspection_report",
            [
                (
                    ["Harbor Chimney Sweeps", "CHIMNEY INSPECTION REPORT"]
                    if k == 1
                    else ["Harbor Chimney Sweeps"],
                    _fill(f"Chimney {k}", 10),
                    [f"Page {k} of 3"],
                )
                for k in (1, 2, 3)
            ],
        )
    )
    # D2 one-page handwritten-style receipt, no letterhead
    docs.append(
        (
            "invoice",
            [
                (
                    [
                        "Thank you for your business!",
                        "Received from owner for gutter cleaning, paid in full.",
                    ],
                    _fill("Receipt", 6),
                    ["Gutter Guys"],
                )
            ],
        )
    )
    # D3 garage door opener manual, 2 pages, unnumbered
    docs.append(
        (
            "manual",
            [
                (["LiftRight 3000 Garage Door Opener", "OWNERS MANUAL"], _fill("Opener", 10), []),
                (["LiftRight 3000 Garage Door Opener"], _fill("Opener", 10), []),
            ],
        )
    )
    # D4, D5 two letters from the same HOA, same phone, different dates
    for d in ("January 5, 2024", "February 9, 2024"):
        docs.append(
            (
                "other",
                [
                    (
                        [
                            "Coastal Bluffs Homeowners Association",
                            "(760) 555-0101 | www.coastalbluffs.example",
                            d,
                        ],
                        _fill("HOA letter", 9),
                        [],
                    )
                ],
            )
        )
    # D6 a neighbour's note, nothing distinctive
    docs.append(
        (
            "other",
            [
                (
                    ["Hi neighbours,", "About the shared fence repair"],
                    _fill("Fence", 8),
                    ["Thanks, Pat"],
                )
            ],
        )
    )
    # D7 a 2-page electrical panel label schedule, then its manufacturer's data sheet
    docs.append(
        (
            "other",
            [
                (["Panel schedule, main service"], _fill("Circuit", 12), []),
                (["Panel schedule, sub panel"], _fill("Circuit", 12), []),
            ],
        )
    )
    docs.append(
        (
            "manual",
            [
                (
                    ["Voltek Breakers Inc.", "Product data sheet"],
                    _fill("Spec", 10),
                    ["www.voltek.example"],
                )
            ],
        )
    )
    return docs


def review_flatten(docs):
    pages, starts, types, p = [], [], [], 1
    for t, ps in docs:
        starts.append(p)
        types.append(t)
        pages += ps
        p += len(ps)
    return pages, starts, types


REVIEW_BUNDLES = {
    "review_b1.pdf": lambda: review_bundle1(True),
    "review_b1b.pdf": lambda: review_bundle1(False),
    "review_b2.pdf": review_bundle2,
    "review_b3.pdf": review_bundle3,
}


def review_layout(name: str) -> tuple[list, list[int], list[str]]:
    """``(pages, document starts, doc types)`` of a review bundle."""
    return review_flatten(REVIEW_BUNDLES[name]())


def _pdf(path: Path, *drawers) -> None:
    from reportlab.pdfgen import canvas as rl_canvas

    # invariant: fixed dates and IDs; pageCompression=0: page streams don't depend
    # on the zlib build (the photo's image stream is still Flate-encoded, so the
    # reproducibility test compares decompressed streams when bytes differ)
    canvas = rl_canvas.Canvas(str(path), pagesize=(PAGE_W, PAGE_H), invariant=1, pageCompression=0)
    canvas.setAuthor("HAL 9000 fixture builder")
    canvas.setTitle(path.stem)
    canvas.setCreator("tests/knowledge/fixtures/build_fixture_corpus.py")
    for draw in drawers:
        draw(canvas)
    canvas.save()


# --------------------------------------------------------------- text files
DOCX_TITLE = "Harborline HP-3000 Heat Pump: Owner's Guide"
DOCX_HEADINGS = ("Filter Maintenance", "Thermostat Settings", "Warranty Registration")


def _fixed_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed timestamps and order, so the DOCX is reproducible."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(2025, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            target.writestr(fixed, source.read(info.filename))
    return out.getvalue()


def write_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    moment = datetime(2025, 1, 1, 0, 0, 0)
    props = document.core_properties
    props.author = "HAL 9000 fixture builder"
    props.last_modified_by = "HAL 9000 fixture builder"
    props.title = DOCX_TITLE
    props.created = moment
    props.modified = moment
    props.revision = 1

    document.add_paragraph(DOCX_TITLE, style="Title")
    document.add_paragraph("Keep this guide with the unit. Model HP-3000, serial HL-000000-00.")

    document.add_heading(DOCX_HEADINGS[0], level=1)
    document.add_paragraph(
        "Replace the 16 x 25 x 1 inch MERV 8 air filter every 90 days, or every 60 days when "
        "pets live in the house. A clogged filter reduces airflow and can freeze the coil."
    )
    document.add_heading("Checking the filter", level=2)
    document.add_paragraph(
        "Slide the filter out of the return grille in the hallway ceiling. Hold it up to a "
        "light; if you cannot see light through it, replace it."
    )

    document.add_heading(DOCX_HEADINGS[1], level=1)
    document.add_paragraph("The factory schedule uses these setpoints:")
    table = document.add_table(rows=3, cols=2)
    for row, (mode, value) in zip(
        table.rows, (("Mode", "Setpoint"), ("Heat", "68 F"), ("Cool", "76 F"))
    ):
        row.cells[0].text = mode
        row.cells[1].text = value
    document.add_paragraph("Hold the MODE button for five seconds to switch to emergency heat.")

    document.add_heading(DOCX_HEADINGS[2], level=1)
    document.add_paragraph(
        "The compressor carries a 10-year limited warranty when the unit is registered within "
        "60 days of installation. Other parts carry a 5-year limited warranty. Registration "
        "number: HL-WR-000123."
    )

    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(_fixed_zip(buffer.getvalue()))


HTML_TITLE = "Irrigation Controller Notes"
HTML_SECTIONS = ("Zones", "Watering Schedule", "Winterizing")

HTML_SOURCE = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{HTML_TITLE}</title>
<style>body {{ font-family: sans-serif; }} .secret {{ display: none; }}</style>
<script>window.trackingPixel = "do-not-index";</script>
</head>
<body>
<h1>{HTML_TITLE}</h1>
<p>Notes for the <b>RainWise 8</b> controller in the garage (a fictional model).</p>
<h2>{HTML_SECTIONS[0]}</h2>
<table>
<tr><th>Zone</th><th>Area</th><th>Heads</th></tr>
<tr><td>1</td><td>Front lawn</td><td>Rotors</td></tr>
<tr><td>2</td><td>Side yard drip</td><td>Drip emitters</td></tr>
<tr><td>3</td><td>Back lawn</td><td>Spray heads</td></tr>
</table>
<h2>{HTML_SECTIONS[1]}</h2>
<ul>
<li>Lawns run Monday, Wednesday and Friday at 5:00 AM for 12 minutes.</li>
<li>Drip runs Tuesday and Saturday for 30 minutes.</li>
</ul>
<p>Rain delay: press <em>RAIN</em> twice for a 48-hour pause.</p>
<script>console.log("not text");</script>
<h2>{HTML_SECTIONS[2]}</h2>
<p>Shut the backflow valve in November, then blow out the lines at 50 psi.<br>
Restart in March.</p>
<noscript>Enable JavaScript.</noscript>
</body>
</html>
"""

MARKDOWN_SECTIONS = ("Opener", "Remote Codes", "Maintenance")
MARKDOWN_SOURCE = f"""# Garage Door Opener

A LiftRight 1/2 HP belt-drive opener (a fictional model), installed in 2019.

## {MARKDOWN_SECTIONS[0]}

The opener is on a dedicated 15-amp circuit, breaker 14 in the garage subpanel.
The travel limits were last adjusted in 2024.

## {MARKDOWN_SECTIONS[1]}

To pair a new remote, press the LEARN button on the motor head once, then press
the remote button within 30 seconds.

```
# this line is inside a code block, not a heading
LEARN -> remote button -> light blinks twice
```

## {MARKDOWN_SECTIONS[2]}

Lubricate the rail and rollers every six months with white lithium grease.
Test the auto-reverse with a 2x4 laid flat under the door every month.
"""

_POOL_SENTENCES = (
    "Checked the pool pump basket and removed leaves",
    "Free chlorine read {c} ppm and pH read {p}",
    "Backwashed the sand filter until the sight glass ran clear",
    "Brushed the walls and the steps, then vacuumed the deep end",
    "Topped up the water level by about {w} inches",
    "The heater was off; the cover stayed on overnight",
    "Added {s} pounds of salt to the chlorinator",
    "Cleaned the skimmer weir and checked the O-ring for cracks",
)


def pool_log_text(target: int = 7000) -> str:
    """Heading-less prose, about ``target`` characters, deterministic."""
    paragraphs: list[str] = []
    day = 0
    while sum(len(p) + 2 for p in paragraphs) < target:
        day += 1
        values = {
            "c": 1 + day % 3,
            "p": f"7.{2 + day % 5}",
            "w": 1 + day % 2,
            "s": 10 * (1 + day % 4),
        }
        sentences = [
            _POOL_SENTENCES[(day + k) % len(_POOL_SENTENCES)].format(**values) + "."
            for k in range(4)
        ]
        paragraphs.append(f"Visit {day}. " + " ".join(sentences))
    return "\n\n".join(paragraphs) + "\n"


# --------------------------------------------------------------------- main
FILE_NAMES = (
    "inspection.pdf",
    "estimate.pdf",
    "invoice.pdf",
    "bundle.pdf",
    "unnumbered_bundle.pdf",
    "mixed_bundle.pdf",
    "review_b1.pdf",
    "review_b1b.pdf",
    "review_b2.pdf",
    "review_b3.pdf",
    "heat_pump_guide.docx",
    "irrigation_notes.html",
    "garage_door.md",
    "pool_log.txt",
)


def build(out_dir: Path = FILES_DIR) -> list[Path]:
    """Write every fixture file into ``out_dir``; returns the paths in :data:`FILE_NAMES` order."""
    out_dir.mkdir(parents=True, exist_ok=True)
    _pdf(out_dir / "inspection.pdf", draw_inspection)
    _pdf(out_dir / "estimate.pdf", draw_estimate)
    _pdf(out_dir / "invoice.pdf", draw_invoice)
    _pdf(out_dir / "bundle.pdf", draw_inspection, draw_estimate, draw_invoice)
    _pdf(out_dir / "unnumbered_bundle.pdf", draw_unnumbered)
    _pdf(out_dir / "mixed_bundle.pdf", draw_mixed)
    for name in REVIEW_BUNDLES:
        pages = review_layout(name)[0]
        _pdf(out_dir / name, lambda canvas, pages=pages: _draw_review(canvas, pages))
    write_docx(out_dir / "heat_pump_guide.docx")
    (out_dir / "irrigation_notes.html").write_text(HTML_SOURCE, encoding="utf-8")
    (out_dir / "garage_door.md").write_text(MARKDOWN_SOURCE, encoding="utf-8")
    (out_dir / "pool_log.txt").write_text(pool_log_text(), encoding="utf-8")
    return [out_dir / name for name in FILE_NAMES]


def main(argv: list[str] | None = None) -> int:  # pragma: no cover - a developer command
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=FILES_DIR, help="output directory")
    args = parser.parse_args(argv)
    for path in build(args.out):
        print(f"{path.stat().st_size:>8}  {path}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
